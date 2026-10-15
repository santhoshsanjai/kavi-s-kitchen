"""Exports for the owner's books: CSV, Excel and PDF of balances, ledger, payments, earnings, deliveries."""
import csv
import io
from datetime import date, datetime, time, timedelta
from enum import Enum
from typing import Any, List, Optional, Sequence, Tuple

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.serialize import jsonable
from app.api.v1.endpoints.admin._common import Admin
from app.core.database import get_db
from app.core.time import BUSINESS_TZ, business_today
from app.models.customer import Customer
from app.models.delivery import DeliveryTask
from app.models.ledger import CustomerLedgerEntry, Payment
from app.models.payout import DriverDaySummary
from app.models.user import User
from app.services.errors import ValidationFailed

router = APIRouter()


class ReportKind(str, Enum):
    balances = "balances"
    ledger = "ledger"
    payments = "payments"
    earnings = "earnings"
    deliveries = "deliveries"


class ReportFormat(str, Enum):
    csv = "csv"
    xlsx = "xlsx"
    pdf = "pdf"


Table = Tuple[str, List[str], List[Sequence[Any]]]


def _bounds(start: date, end: date) -> Tuple[datetime, datetime]:
    return (datetime.combine(start, time.min, BUSINESS_TZ), datetime.combine(end + timedelta(days=1), time.min, BUSINESS_TZ))


def _local(ts: datetime) -> str:
    return ts.astimezone(BUSINESS_TZ).strftime("%Y-%m-%d %H:%M")


async def _build(db: AsyncSession, kind: ReportKind, start: date, end: date) -> Table:
    lo, hi = _bounds(start, end)
    if kind == ReportKind.balances:
        rows = (await db.execute(
            select(Customer.name, Customer.phone, func.sum(CustomerLedgerEntry.amount))
            .join(CustomerLedgerEntry, CustomerLedgerEntry.customer_id == Customer.id)
            .where(CustomerLedgerEntry.created_at < hi).group_by(Customer.id).order_by(Customer.name))).all()
        return (f"Customer balances as of {end}", ["Customer", "Phone", "Balance (Rs)"],
                [(n, p, float(b)) for n, p, b in rows])
    if kind == ReportKind.ledger:
        rows = (await db.execute(
            select(CustomerLedgerEntry, Customer.name).join(Customer, Customer.id == CustomerLedgerEntry.customer_id)
            .where(CustomerLedgerEntry.created_at >= lo, CustomerLedgerEntry.created_at < hi)
            .order_by(CustomerLedgerEntry.created_at))).all()
        return (f"Ledger {start} to {end}", ["Date", "Customer", "Type", "Description", "Amount (Rs)", "Reason"],
                [(_local(e.created_at), n, e.entry_type.value, e.description, float(e.amount), e.reason or "")
                 for e, n in rows])
    if kind == ReportKind.payments:
        rows = (await db.execute(
            select(Payment, Customer.name).join(Customer, Customer.id == Payment.customer_id)
            .where(Payment.received_at >= lo, Payment.received_at < hi).order_by(Payment.received_at))).all()
        return (f"Payments {start} to {end}", ["Received", "Customer", "Method", "Amount (Rs)", "Reference"],
                [(_local(p.received_at), n, p.method.value, float(p.amount), p.reference or "") for p, n in rows])
    if kind == ReportKind.earnings:
        rows = (await db.execute(
            select(DriverDaySummary, User.full_name, User.driver_code).join(User, User.id == DriverDaySummary.driver_id)
            .where(DriverDaySummary.service_date.between(start, end))
            .order_by(DriverDaySummary.service_date, User.full_name))).all()
        return (f"Driver earnings {start} to {end}", ["Date", "Driver", "Code", "Status", "Total (Rs)"],
                [(s.service_date.isoformat(), n, c or "", s.status.value, float(s.total_amount)) for s, n, c in rows])
    rows = (await db.execute(
        select(DeliveryTask, Customer.name, User.full_name).join(Customer, Customer.id == DeliveryTask.customer_id)
        .outerjoin(User, User.id == DeliveryTask.assigned_driver_id)
        .where(DeliveryTask.service_date.between(start, end))
        .order_by(DeliveryTask.service_date, DeliveryTask.meal_window, DeliveryTask.stop_sequence))).all()
    return (f"Deliveries {start} to {end}", ["Date", "Meal", "Customer", "Driver", "Status", "Delivery fee (Rs)", "Note"],
            [(t.service_date.isoformat(), t.meal_window.value, n, d or "", t.status.value,
              float(t.customer_delivery_charge or 0), t.failure_reason or t.delivery_note or "") for t, n, d in rows])


def _csv(table: Table) -> bytes:
    _, header, rows = table
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens it as UTF-8


def _xlsx(table: Table) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    title, header, rows = table
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(header)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col[2:]) if len(col) > 2 else 10
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(table: Table) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table as PdfTable, TableStyle
    title, header, rows = table
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=title)
    styles = getSampleStyleSheet()
    body = [[Paragraph(str(c), styles["BodyText"]) for c in r] for r in rows]
    grid = PdfTable([header, *body], repeatRows=1)
    grid.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14532d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    doc.build([Paragraph(f"Kavis Kitchen: {title}", styles["Title"]), Spacer(1, 8),
               grid if rows else Paragraph("No records in this period.", styles["BodyText"])])
    return buf.getvalue()


MEDIA = {
    ReportFormat.csv: "text/csv",
    ReportFormat.xlsx: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ReportFormat.pdf: "application/pdf",
}


@router.get("/reports/{kind}")
async def export_report(kind: ReportKind, format: ReportFormat = ReportFormat.csv, start: Optional[date] = None,
                        end: Optional[date] = None, _: dict = Admin, db: AsyncSession = Depends(get_db)):
    end = end or business_today()
    start = start or end.replace(day=1)
    if start > end:
        raise ValidationFailed("Start date must be on or before the end date")
    if (end - start).days > 366:
        raise ValidationFailed("Reports cover at most one year")
    table = await _build(db, kind, start, end)
    data = {ReportFormat.csv: _csv, ReportFormat.xlsx: _xlsx, ReportFormat.pdf: _pdf}[format](table)
    filename = f"kavis-{kind.value}-{start}-{end}.{format.value}"
    return StreamingResponse(io.BytesIO(data), media_type=MEDIA[format],
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/reports/{kind}/preview")
async def preview_report(kind: ReportKind, start: Optional[date] = None, end: Optional[date] = None,
                         _: dict = Admin, db: AsyncSession = Depends(get_db)):
    end = end or business_today()
    start = start or end.replace(day=1)
    title, header, rows = await _build(db, kind, start, end)
    return {"title": title, "columns": header, "rows": jsonable([list(r) for r in rows[:200]]), "total_rows": len(rows)}
